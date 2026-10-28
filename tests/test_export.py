"""X-1: downloadable notes (PLAN M6c): Markdown, Word and PDF carry the same notebook as the page, in the same order,
with the answers at the end; model-written text cannot break any format."""

import io
import json
from pathlib import Path

import pytest

from dashboard import export

EXAMPLES = Path(__file__).resolve().parents[1] / "schemas" / "examples"
TRICKY = 'Use a[i](x) -- not **this** or __init__; C:\\path\\to # 50% < 3 > 1 “quoted” – ≥ → ✓ 😀'


def notebook(finished=True) -> dict:
    """Every element kind once (from the consolidation fixture), a fallback topic, drafts, glossary, info."""
    done = json.loads((EXAMPLES / "consolidation" / "valid-llm.json").read_text(encoding="utf-8"))
    els = [{k: v for k, v in e.items() if k != "topic"} for e in done["elements"]]
    example = {"kind": "example", "t": 640.0, "said": "", "title": "Picking a random responder",
               "sentences": ["random.randint(16, 272) returned 75."], "steps": ["Import random", "Call randint"],
               "code": "import random\nrandom.randint(16, 272)"}
    tricky = {"kind": "para", "t": 700.0, "said": "", "lead": "Edge cases", "sentences": [TRICKY], "terms": [],
              "background": "Not said: " + TRICKY}
    return {
        "topics": [
            {"n": 1, "title": "Part 1", "start": 0.0, "end": 360.0, "question": None, "answer": None,
             "generator": "fallback", "elements": [{"kind": "para", "t": 130.0, "said": "", "lead": "",
                                                     "sentences": ["Today is about computation."], "terms": [],
                                                     "background": None}]},
            {"n": 2, "title": "Imperative knowledge", "start": 360.0, "end": 750.0,
             "question": "What is imperative knowledge?", "answer": "A recipe: steps that say how to do something.",
             "generator": "llm", "elements": els + [example, tricky],
             "additions": [{"kind": "steps", "t": 640.0, "said": "", "lead": "Your step-by-step solution",
                            "items": ["Pick a number", "Find the student"]}]},
        ],
        "drafts": [] if finished else [{"kind": "para", "t": 760.0, "said": "", "lead": "Still live",
                                        "sentences": ["A draft sentence."], "terms": [], "background": None,
                                        "sentence_k": [8]}],
        "latest_k": 8, "since": 750.0, "reach": 840.0 if not finished else 750.0, "said_until": 850.0,
        "now_text": "", "glossary": [{"term": "Imperative knowledge", "definition": els[0]["terms"][0]["definition"]}],
        "info": [{"text": "Problem set 0 is due on Thursday.", "t": 640.0}], "updates": 9, "finished": finished,
        "status": "",
    }


def test_markdown_follows_the_page_and_puts_the_answers_at_the_end():
    md = export.to_markdown(notebook(), "Python · Lecture 1")
    order = ["# Python · Lecture 1", "## Contents", "## 1. Part 1", "Written without the language model",
             "## 2. Imperative knowledge", "**Imperative knowledge.** **Imperative knowledge** is a recipe",
             "| Kind | What it is |", "**The recipe the students gave:**", "**Remember:**", "> **Example · Picking",
             "### Your additions", "**Check yourself:** What is imperative knowledge?", "## Glossary",
             "## Answers to “Check yourself”", "2. *What is imperative knowledge?* A recipe", "## Course info"]
    positions = [md.index(x) for x in order]
    assert positions == sorted(positions)
    assert "*[06:50]*" in md and "Background (not from the lecture)" in md
    assert "Live notes since" not in md                                          # finished: no drafts


def test_markdown_escapes_model_text():
    md = export.to_markdown(notebook(), "x")
    assert r"a\[i\](x) -- not \*\*this\*\* or \_\_init\_\_" in md and r"\# 50% \< 3 \> 1" in md


def test_a_live_lecture_exports_its_drafts_and_says_how_far_the_notes_reach():
    md = export.to_markdown(notebook(finished=False), "x")
    assert "Live lecture: organised notes up to 12:30, draft notes up to 14:00" in md
    assert "## Live notes since 12:30 (draft)" in md and "A draft sentence." in md


def test_word_has_the_headings_tables_and_runs():
    from docx import Document
    doc = Document(io.BytesIO(export.to_docx(notebook(), "Python · Lecture 1")))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith(("Heading", "Title"))]
    assert headings[:4] == ["Python · Lecture 1", "Contents", "1. Part 1", "2. Imperative knowledge"]
    assert "Your additions" in headings and headings[-3:] == ["Glossary", "Answers to “Check yourself”",
                                                              "Course info"]
    para = next(p for p in doc.paragraphs if p.text.startswith("Imperative knowledge. Imperative"))
    assert [r.text for r in para.runs if r.bold][:2] == ["Imperative knowledge. ", "Imperative knowledge"]
    assert [[c.text for c in row.cells] for row in doc.tables[0].rows][0] == ["Kind", "What it is"]
    assert any(TRICKY in p.text for p in doc.paragraphs)                        # text kept as is


@pytest.mark.parametrize("finished", [True, False])
def test_pdf_renders_every_element_and_any_text(finished):
    pdf = export.to_pdf(notebook(finished), "Python · Lecture 1 – “notes”")
    assert pdf.startswith(b"%PDF") and pdf.rstrip().endswith(b"%%EOF") and len(pdf) > 10_000


def test_formats_table_drives_the_download_buttons():
    assert list(export.FORMATS) == ["Markdown", "Word", "PDF"]
    for fmt, (ext, mime, build) in export.FORMATS.items():
        data = build(notebook(), "t")
        assert isinstance(data, bytes) and data and ext in ("md", "docx", "pdf")
    assert export.file_stem("Python · Lecture 1 (Bell)") == "Python_Lecture_1_Bell"
